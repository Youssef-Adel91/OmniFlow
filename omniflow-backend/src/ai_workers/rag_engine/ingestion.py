"""
ai_workers/rag_engine/ingestion.py — Document text extraction & chunking

Pure, synchronous, dependency-light helpers shared by the Celery ingestion task
(`omniflow.ingest_knowledge_document`). No I/O, no DB, no network — everything
here takes bytes and returns strings, which makes it trivially unit-testable.

PARSER CHOICE
-------------
`unstructured[all-docs]` is declared in pyproject but is NOT actually importable
in the current environment (only `unstructured_client` / `unstructured_pytesseract`
are installed, and the full extra drags in torch + poppler + tesseract system
packages). So the strategy is:

    1. Try `unstructured.partition.auto.partition` if it imports — best quality,
       handles tables and scanned PDFs.
    2. Otherwise fall back to the lightweight parsers that ARE installed:
         pdf  → pypdf            (pypdf 6.x)
         docx → python-docx      (paragraphs + table cells)
         txt  → utf-8 / cp1256 decode
         csv  → csv module, rows flattened to "col: value" lines

The fallback is the path that actually runs today. It is text-layer only: a
scanned/image-only PDF yields no text and the task marks the document `failed`
with a clear message rather than indexing an empty document.

CHUNKING
--------
Character-based, not token-based, deliberately: a tokenizer dependency would be
a third embedding-model-specific moving part for very little gain at this size.

    _CHUNK_CHARS   = 2400  ≈ 600 tokens of Arabic prose
    _CHUNK_OVERLAP = 300   ≈ 75 tokens, so a fact spanning a boundary survives

Splitting prefers paragraph breaks, then sentence enders (including the Arabic
full stop and question mark), then falls back to a hard character cut.
"""
from __future__ import annotations

import csv
import io
import re
from typing import Final

import structlog

logger = structlog.get_logger(__name__)

# Supported upload types — kept in sync with the router's validation table.
SUPPORTED_FILE_TYPES: Final[frozenset[str]] = frozenset({"pdf", "docx", "txt", "csv"})

_CHUNK_CHARS: Final[int] = 2400
_CHUNK_OVERLAP: Final[int] = 300

# Minimum characters of extracted text before we consider a document usable.
_MIN_USABLE_CHARS: Final[int] = 30

# Sentence terminators: latin + Arabic full stop/question mark, and newlines.
_BREAK_CHARS: Final[str] = ".!?؟۔\n"


class ExtractionError(RuntimeError):
    """Raised when a document cannot be parsed into usable text."""


# ══════════════════════════════════════════════════════════════════════════════
# Text extraction
# ══════════════════════════════════════════════════════════════════════════════

def extract_text(file_bytes: bytes, file_type: str, *, filename: str = "") -> str:
    """
    Extract plain text from an uploaded document.

    Args:
        file_bytes — Raw object body fetched from S3/MinIO.
        file_type  — One of SUPPORTED_FILE_TYPES (lower-case, no dot).
        filename   — Original filename, used only for logging / unstructured hints.

    Returns:
        Normalised plain text (whitespace collapsed, no NUL bytes).

    Raises:
        ExtractionError — unsupported type, corrupt file, or no text layer.
    """
    ftype = (file_type or "").lower().lstrip(".")
    if ftype not in SUPPORTED_FILE_TYPES:
        raise ExtractionError(f"Unsupported file type: {file_type!r}")

    text = _try_unstructured(file_bytes, filename)

    if text is None:
        if ftype == "pdf":
            text = _extract_pdf(file_bytes)
        elif ftype == "docx":
            text = _extract_docx(file_bytes)
        elif ftype == "csv":
            text = _extract_csv(file_bytes)
        else:  # txt
            text = _decode(file_bytes)

    text = _normalise(text)

    if len(text) < _MIN_USABLE_CHARS:
        raise ExtractionError(
            "No readable text found in the document. If this is a scanned PDF, "
            "run OCR on it first or upload a text-based version."
        )
    return text


def _try_unstructured(file_bytes: bytes, filename: str) -> str | None:
    """Best-effort `unstructured` parse. Returns None if unavailable/failed."""
    try:
        from unstructured.partition.auto import partition  # noqa: PLC0415
    except Exception:  # ImportError, or a missing system dep at import time
        return None

    try:
        elements = partition(file=io.BytesIO(file_bytes), metadata_filename=filename or None)
        text = "\n\n".join(str(el) for el in elements if str(el).strip())
        return text or None
    except Exception as exc:  # noqa: BLE001 — never let the good path break the fallback
        logger.warning("unstructured_partition_failed", error=str(exc)[:300])
        return None


def _extract_pdf(file_bytes: bytes) -> str:
    try:
        from pypdf import PdfReader  # noqa: PLC0415
    except ImportError as exc:  # pragma: no cover
        raise ExtractionError("pypdf is not installed; cannot parse PDF.") from exc

    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        if getattr(reader, "is_encrypted", False):
            # Try the empty password — many "protected" PDFs use one.
            try:
                reader.decrypt("")
            except Exception as exc:  # noqa: BLE001
                raise ExtractionError("PDF is password-protected.") from exc
        pages = [(page.extract_text() or "") for page in reader.pages]
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"Corrupt or unreadable PDF: {exc}") from exc

    return "\n\n".join(p for p in pages if p.strip())


def _extract_docx(file_bytes: bytes) -> str:
    try:
        import docx  # noqa: PLC0415  (python-docx)
    except ImportError as exc:  # pragma: no cover
        raise ExtractionError("python-docx is not installed; cannot parse DOCX.") from exc

    try:
        document = docx.Document(io.BytesIO(file_bytes))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"Corrupt or unreadable DOCX: {exc}") from exc

    parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]

    # Tables carry most of the value in price lists / catalogues.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def _extract_csv(file_bytes: bytes) -> str:
    raw = _decode(file_bytes)
    try:
        reader = csv.reader(io.StringIO(raw))
        rows = list(reader)
    except Exception:  # noqa: BLE001 — fall back to the raw text
        return raw

    if not rows:
        return ""

    header = [h.strip() for h in rows[0]]
    lines: list[str] = []
    for row in rows[1:]:
        pairs = [
            f"{header[i] if i < len(header) else f'col{i}'}: {value.strip()}"
            for i, value in enumerate(row)
            if value and value.strip()
        ]
        if pairs:
            lines.append(" | ".join(pairs))

    # No data rows → keep the header so the doc is not silently empty.
    return "\n".join(lines) if lines else " | ".join(header)


def _decode(file_bytes: bytes) -> str:
    """Decode bytes as UTF-8, falling back to Windows-1256 (legacy Arabic)."""
    for encoding in ("utf-8", "utf-8-sig", "cp1256", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="replace")


def _normalise(text: str) -> str:
    """Strip NULs, collapse runs of blank lines and trailing spaces."""
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ══════════════════════════════════════════════════════════════════════════════
# Chunking
# ══════════════════════════════════════════════════════════════════════════════

def chunk_text(
    text: str,
    *,
    chunk_chars: int = _CHUNK_CHARS,
    overlap: int = _CHUNK_OVERLAP,
) -> list[str]:
    """
    Split text into overlapping chunks on natural boundaries.

    Guarantees:
      * every returned chunk is non-empty after stripping
      * consecutive chunks overlap by up to `overlap` characters
      * the function always terminates (the cursor advances by >= 1 char)

    Returns [] for empty input.
    """
    text = (text or "").strip()
    if not text:
        return []
    if len(text) <= chunk_chars:
        return [text]

    overlap = max(0, min(overlap, chunk_chars // 2))

    chunks: list[str] = []
    start = 0
    length = len(text)

    while start < length:
        end = min(start + chunk_chars, length)

        if end < length:
            split_at = _find_break(text, start, end)
            if split_at > start:
                end = split_at

        piece = text[start:end].strip()
        if piece:
            chunks.append(piece)

        if end >= length:
            break

        next_start = end - overlap
        start = next_start if next_start > start else end  # always make progress

    return chunks


def _find_break(text: str, start: int, end: int) -> int:
    """
    Find the best split point at or before `end`, no earlier than 60% into the
    window (so we never produce a tiny chunk just because a period appeared
    early). Returns `end` unchanged when nothing better exists.
    """
    floor = start + int((end - start) * 0.6)

    paragraph = text.rfind("\n\n", floor, end)
    if paragraph != -1:
        return paragraph + 2

    for index in range(end - 1, floor, -1):
        if text[index] in _BREAK_CHARS:
            return index + 1

    return end


__all__ = [
    "SUPPORTED_FILE_TYPES",
    "ExtractionError",
    "chunk_text",
    "extract_text",
]
