"""
File sniffing + tabular readers for bulk imports.

The file type is decided from the *bytes* (magic numbers / container contents),
never from the extension alone, so a renamed executable or a `.csv` that is
really a PDF cannot slip through. Readers are bounded (rows, columns, zip
expansion) so a hostile file cannot exhaust memory.
"""
from __future__ import annotations

import csv
import io
import json
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Literal

Kind = Literal["csv", "tsv", "xlsx", "xls", "json", "pdf", "docx", "image"]

TABULAR_KINDS = ("csv", "tsv", "xlsx", "xls", "json")
LLM_KINDS = ("pdf", "docx", "image")

MAX_COLUMNS = 200
MAX_ZIP_EXPANDED_BYTES = 300 * 1024 * 1024     # xlsx/docx are zips: refuse zip bombs
MAX_ZIP_RATIO = 200
_OLE_MAGIC = b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1"


class ImportFileError(ValueError):
    """The upload cannot be used. `code` is stable for the UI/tests; message is Arabic."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass
class Table:
    headers: list[str]
    rows: list[list[str]]
    # Set when the rows came from an LLM (PDF/DOCX/image): the user must review them.
    extracted_by_llm: bool = False
    notes: list[str] = field(default_factory=list)


# ── sniffing ─────────────────────────────────────────────────────────────────

def _ext(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _zip_members(data: bytes) -> list[zipfile.ZipInfo]:
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ImportFileError("corrupt_file", "الملف تالف أو ليس بصيغة صالحة.") from exc
    infos = zf.infolist()
    total = sum(i.file_size for i in infos)
    compressed = max(1, sum(i.compress_size for i in infos))
    if total > MAX_ZIP_EXPANDED_BYTES or total / compressed > MAX_ZIP_RATIO:
        raise ImportFileError("zip_bomb", "الملف مرفوض: حجمه بعد فك الضغط كبير بشكل غير طبيعي.")
    return infos


def detect_kind(data: bytes, filename: str) -> Kind:
    """Decide the file kind from its content; reject anything unsupported or mismatched."""
    ext = _ext(filename)
    head = data[:16]
    if not data:
        raise ImportFileError("empty_file", "الملف فارغ.")

    if head.startswith(b"%PDF-"):
        kind: Kind = "pdf"
    elif head.startswith(b"PK\x03\x04"):
        names = {i.filename for i in _zip_members(data)}
        if "xl/workbook.xml" in names:
            kind = "xlsx"
        elif "word/document.xml" in names:
            kind = "docx"
        else:
            raise ImportFileError("unsupported_type", "صيغة الملف غير مدعومة (أرشيف ZIP غير معروف).")
    elif head.startswith(_OLE_MAGIC):
        if ext != "xls":
            raise ImportFileError("unsupported_type", "صيغة Office القديمة غير مدعومة إلا بصيغة XLS.")
        kind = "xls"
    elif head.startswith(b"\x89PNG\r\n\x1a\n") or head.startswith(b"\xff\xd8\xff") or (
        head[:4] == b"RIFF" and data[8:12] == b"WEBP"
    ):
        kind = "image"
    else:
        sample = data[:8192]
        utf16 = data.startswith((b"\xff\xfe", b"\xfe\xff"))     # UTF-16 text is full of NULs but is text
        if b"\x00" in sample and not utf16:
            raise ImportFileError("unsupported_type", "صيغة الملف غير مدعومة (ملف ثنائي غير معروف).")
        stripped = b"" if utf16 else data.lstrip(b"\xef\xbb\xbf \t\r\n")[:1]
        if ext == "json" or stripped in (b"{", b"["):
            kind = "json"
        elif ext == "tsv" or (not utf16 and ext not in ("csv", "txt") and sample.count(b"\t") > sample.count(b",")):
            kind = "tsv"
        else:
            kind = "csv"

    expected = {"csv": {"csv", "txt", "tsv"}, "tsv": {"tsv", "txt", "csv"}, "json": {"json", "txt"},
                "xlsx": {"xlsx", "xlsm"}, "xls": {"xls"}, "docx": {"docx"}, "pdf": {"pdf"},
                "image": {"png", "jpg", "jpeg", "webp"}}[kind]
    if ext and ext not in expected:
        raise ImportFileError(
            "extension_mismatch", f"امتداد الملف (.{ext}) لا يطابق محتواه الفعلي ({kind}).",
        )
    return kind


# ── readers ──────────────────────────────────────────────────────────────────

def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() and abs(value) < 1e15 else repr(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    s = str(value)
    return "".join(ch for ch in s if ch == "\t" or ch == "\n" or ch == "\r" or ord(ch) >= 32).strip()


def _finish(headers: list[str], rows: list[list[str]], max_rows: int) -> Table:
    """Common post-processing: pad rows, name blank/duplicate headers, drop empty rows/columns."""
    width = min(MAX_COLUMNS, max([len(headers), *(len(r) for r in rows)] or [0]))
    headers = (headers + [""] * width)[:width]
    seen: dict[str, int] = {}
    named: list[str] = []
    for i, h in enumerate(headers):
        h = h.strip() or f"عمود {i + 1}"
        n = seen.get(h, 0)
        seen[h] = n + 1
        named.append(h if n == 0 else f"{h} ({n + 1})")
    rows = [(r + [""] * width)[:width] for r in rows if any(c.strip() for c in r)]
    # drop trailing columns that are blank in the header and every row
    while named and named[-1].startswith("عمود ") and not any(r[len(named) - 1] for r in rows):
        named.pop()
        rows = [r[: len(named)] for r in rows]
    if not named:
        raise ImportFileError("no_columns", "لم يتم العثور على أعمدة في الملف.")
    if not rows:
        raise ImportFileError("no_rows", "الملف لا يحتوي على صفوف بيانات.")
    if len(rows) > max_rows:
        raise ImportFileError("too_many_rows", f"عدد الصفوف ({len(rows):,}) يتجاوز الحد المسموح ({max_rows:,}).")
    return Table(headers=named, rows=rows)


def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    try:
        from charset_normalizer import from_bytes

        best = from_bytes(data).best()
    except Exception:  # noqa: BLE001 - detection is best-effort
        best = None
    if best is not None:
        return str(best)
    return data.decode("cp1256", errors="replace")   # common legacy Arabic Windows encoding


def _read_delimited(data: bytes, kind: Kind, max_rows: int) -> Table:
    text = _decode(data)
    if kind == "tsv":
        delimiter = "\t"
    else:
        first = text.split("\n", 1)[0]
        delimiter = max([",", ";", "\t", "|"], key=first.count)
        if first.count(delimiter) == 0:
            delimiter = ","
    csv.field_size_limit(1_000_000)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    out: list[list[str]] = []
    for i, row in enumerate(reader):
        out.append([_cell(c) for c in row])
        if i > max_rows + 5_000:   # blank-heavy files: stop reading long after the cap
            break
    if not out:
        raise ImportFileError("no_rows", "الملف فارغ.")
    return _finish(out[0], out[1:], max_rows)


def _read_xlsx(data: bytes, max_rows: int) -> Table:
    from openpyxl import load_workbook

    _zip_members(data)
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ImportFileError("corrupt_file", "تعذّر قراءة ملف Excel (تالف أو محمي بكلمة مرور).") from exc
    try:
        for ws in wb.worksheets:
            rows: list[list[str]] = []
            for i, raw in enumerate(ws.iter_rows(values_only=True)):
                cells = [_cell(c) for c in raw[:MAX_COLUMNS]]
                if any(cells) or rows:
                    rows.append(cells)
                if i > max_rows + 5_000:
                    break
            if len(rows) >= 2:
                return _finish(rows[0], rows[1:], max_rows)
    finally:
        wb.close()
    raise ImportFileError("no_rows", "لم يتم العثور على بيانات في أي ورقة من ملف Excel.")


def _read_xls(data: bytes, max_rows: int) -> Table:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
    except Exception as exc:  # noqa: BLE001
        raise ImportFileError("corrupt_file", "تعذّر قراءة ملف XLS (تالف أو محمي).") from exc
    for sheet in book.sheets():
        if sheet.nrows >= 2:
            rows = [[_cell(sheet.cell_value(r, c)) for c in range(min(sheet.ncols, MAX_COLUMNS))]
                    for r in range(min(sheet.nrows, max_rows + 5_001))]
            return _finish(rows[0], rows[1:], max_rows)
    raise ImportFileError("no_rows", "لم يتم العثور على بيانات في ملف XLS.")


def _read_json(data: bytes, max_rows: int) -> Table:
    try:
        doc = json.loads(_decode(data))
    except json.JSONDecodeError as exc:
        raise ImportFileError("invalid_json", f"ملف JSON غير صالح (السطر {exc.lineno}).") from exc
    if isinstance(doc, dict):
        for key in ("items", "data", "rows", "properties", "listings", "customers", "results"):
            if isinstance(doc.get(key), list):
                doc = doc[key]
                break
        else:
            doc = [doc]
    if not isinstance(doc, list) or not all(isinstance(x, dict) for x in doc):
        raise ImportFileError("invalid_json", "يجب أن يكون JSON مصفوفة من الكائنات (كل كائن = صف).")
    headers: list[str] = []
    for obj in doc[: max_rows + 1]:
        for k in obj:
            if k not in headers and len(headers) < MAX_COLUMNS:
                headers.append(str(k))
    rows = [[_cell(json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v)
             for v in (obj.get(h) for h in headers)] for obj in doc]
    return _finish(headers, rows, max_rows)


def read_tabular(data: bytes, kind: Kind, max_rows: int) -> Table:
    if kind in ("csv", "tsv"):
        return _read_delimited(data, kind, max_rows)
    if kind == "xlsx":
        return _read_xlsx(data, max_rows)
    if kind == "xls":
        return _read_xls(data, max_rows)
    if kind == "json":
        return _read_json(data, max_rows)
    raise ImportFileError("unsupported_type", "هذه الصيغة تحتاج استخراجًا بالذكاء الاصطناعي.")


def read_docx_table(data: bytes, max_rows: int) -> Table | None:
    """Largest real table inside a .docx, or None when the document is prose."""
    import docx

    _zip_members(data)
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ImportFileError("corrupt_file", "تعذّر قراءة ملف Word.") from exc
    best: list[list[str]] = []
    for t in document.tables:
        rows = [[_cell(c.text) for c in r.cells] for r in t.rows]
        if len(rows) > len(best):
            best = rows
    if len(best) >= 2 and len(best[0]) >= 2:
        return _finish(best[0], best[1:], max_rows)
    return None


def docx_text(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs if p.text.strip())


def pdf_text(data: bytes, max_pages: int = 60) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ImportFileError("encrypted_pdf", "ملف PDF محمي بكلمة مرور.")
        return "\n".join((p.extract_text() or "") for p in reader.pages[:max_pages])
    except ImportFileError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ImportFileError("corrupt_file", "تعذّر قراءة ملف PDF.") from exc
