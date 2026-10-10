"""Customer bulk import: upload -> dry-run -> background commit (same job machinery as the property import).

Reuses from `property_import`: private object storage, the DB-backed job row with lease/checkpoint/resume,
the error-report CSV, progress serialisation, one-active-job-per-tenant and the audit helper. Only the
schema (`bulk_import/customer_schema.py`) and the row writer below are customer specific.

Identity: the natural key is the normalised phone (E.164). A row whose phone already belongs to a customer
(`phone_e164` or the messaging contact key) is merged into that customer according to `on_duplicate`.

PDPL consent: imported customers are NOT opted in unless the person importing ticks the explicit declaration
("I hold these customers' consent to contact them") -- `options.consent_confirmed`. Even then a customer who
opted out earlier is never re-opted-in, and the declaration is audit-logged (who, when, how many).
"""
from __future__ import annotations

import asyncio
import io
import uuid
from datetime import datetime, timedelta, timezone
from typing import Any

import structlog
from sqlalchemy import select, update
from sqlalchemy.dialects.postgresql import insert as pg_insert

from src.shared.core.config import get_settings
from src.shared.db.models import Customer, ImportJob
from src.shared.db.session import get_tenant_session
from src.shared.services import customers as customer_svc
from src.shared.services import property_import as base
from src.shared.services.bulk_import import llm_extract
from src.shared.services.bulk_import.customer_schema import CUSTOMER_SCHEMA, MAX_TAG_LEN, MAX_TAGS
from src.shared.services.bulk_import.files import LLM_KINDS, ImportFileError, detect_kind, read_tabular
from src.shared.services.bulk_import.mapping import suggest_mapping, validate_mapping
from src.shared.services.bulk_import.validate import RowResult, validate_table

logger = structlog.get_logger(__name__)

KIND = "customers"
ON_DUPLICATE = ("skip", "update")
CONSENT_FIELD = "marketing_opt_in"


# ── 1. upload ─────────────────────────────────────────────────────────────────

async def create_upload(session, *, tenant_id: uuid.UUID, user_id: uuid.UUID | None, filename: str, data: bytes) -> dict[str, Any]:
    cfg = get_settings()
    if len(data) > cfg.import_max_file_mb * 1024 * 1024:
        raise ImportFileError("file_too_large", f"حجم الملف يتجاوز الحد المسموح ({cfg.import_max_file_mb} ميجابايت).")
    filename = base._safe_filename(filename)
    kind = detect_kind(data, filename)
    if kind in LLM_KINDS:
        table = await llm_extract.extract_table(kind, data, CUSTOMER_SCHEMA)
    else:
        table = read_tabular(data, kind, cfg.import_max_rows)

    suggestions = suggest_mapping(table.headers, CUSTOMER_SCHEMA)
    unmapped_required = [s.field for s in suggestions if not s.header and CUSTOMER_SCHEMA.by_name()[s.field].required]
    if unmapped_required and not table.extracted_by_llm:
        taken = {s.header for s in suggestions if s.header}
        refined = await llm_extract.refine_mapping(table.headers, table.rows, unmapped_required, taken, CUSTOMER_SCHEMA)
        for s in suggestions:
            if s.field in refined:
                s.header, s.confidence, s.source = refined[s.field], 0.6, "llm"

    import_id = uuid.uuid4()
    keys = base._keys(tenant_id, import_id)
    await base.storage.put(keys["source"], data, "application/octet-stream")
    await base.storage.put(keys["rows"], base._table_to_json(table), "application/json")
    session.add(ImportJob(
        import_id=import_id, tenant_id=tenant_id, kind=KIND, status="uploaded", filename=filename, file_kind=kind,
        file_size=len(data), source_key=keys["source"], rows_key=keys["rows"], columns=table.headers,
        total_rows=len(table.rows), created_by=user_id))
    await session.flush()
    logger.info("customer_import_uploaded", import_id=str(import_id), tenant_id=str(tenant_id), kind=kind, rows=len(table.rows))
    return {
        "import_id": str(import_id), "file_kind": kind, "filename": filename, "total_rows": len(table.rows),
        "detected_columns": table.headers, "suggested_mapping": [s.as_dict() for s in suggestions],
        "fields": [{"name": f.name, "label": f.label_ar, "required": f.required, "example": f.example} for f in CUSTOMER_SCHEMA.fields],
        "sample_rows": [dict(zip(table.headers, r)) for r in table.rows[:5]],
        "extracted_by_llm": table.extracted_by_llm, "notes": table.notes,
        "limits": {"max_file_mb": cfg.import_max_file_mb, "max_rows": cfg.import_max_rows},
    }


# ── 2. validate (dry run) ─────────────────────────────────────────────────────

def normalize_options(options: dict[str, Any] | None) -> dict[str, Any]:
    options = options or {}
    on_dup = options.get("on_duplicate", "skip")
    if on_dup not in ON_DUPLICATE:
        raise ValueError(f"on_duplicate يجب أن يكون أحد: {', '.join(ON_DUPLICATE)}")
    raw_defaults = {k: v for k, v in (options.get("defaults") or {}).items() if v not in (None, "")}
    unknown = set(raw_defaults) - {"city"}
    if unknown:
        raise ValueError(f"قيم افتراضية غير مدعومة: {', '.join(sorted(unknown))}")
    tags = [str(t).strip() for t in (options.get("tags") or []) if str(t).strip()]
    tags = list(dict.fromkeys(tags))
    if len(tags) > MAX_TAGS or any(len(t) > MAX_TAG_LEN for t in tags):
        raise ValueError(f"الوسوم: حتى {MAX_TAGS} وسمًا وكل وسم حتى {MAX_TAG_LEN} حرفًا")
    return {
        "on_duplicate": on_dup,
        "defaults": {k: str(v)[:100] for k, v in raw_defaults.items()},
        "tags": tags,
        "consent_confirmed": bool(options.get("consent_confirmed", False)),
        "opt_in_all": bool(options.get("opt_in_all", False)),
    }


async def _existing_by_phone(session, tenant_id: uuid.UUID, phones: list[str]) -> dict[str, Customer]:
    found: dict[str, Customer] = {}
    for i in range(0, len(phones), 1000):
        chunk = phones[i:i + 1000]
        rows = (await session.execute(select(Customer).where(
            Customer.tenant_id == tenant_id, (Customer.phone_e164.in_(chunk)) | (Customer.unified_phone.in_(chunk))))).scalars().all()
        for c in rows:
            for key in (c.phone_e164, c.unified_phone):
                if key in chunk:
                    found.setdefault(key, c)
    return found


def _wants_consent(values: dict[str, Any], options: dict[str, Any]) -> bool:
    return bool(values.get(CONSENT_FIELD)) or (options["opt_in_all"] and values.get(CONSENT_FIELD) is None)


async def validate_job(session, job: ImportJob, mapping_in: dict[str, str | None], options_in: dict[str, Any] | None) -> dict[str, Any]:
    if job.status in base.ACTIVE or job.status in base.TERMINAL:
        raise ValueError("لا يمكن إعادة التحقق من عملية استيراد بدأت أو انتهت.")
    table = base._table_from_json(await base.storage.get(job.rows_key))
    mapping = validate_mapping(mapping_in, table.headers, CUSTOMER_SCHEMA)
    options = normalize_options(options_in)

    required_missing = [f.label_ar for f in CUSTOMER_SCHEMA.fields if f.required and f.name not in mapping]
    results = list(validate_table(table, mapping, CUSTOMER_SCHEMA, options["defaults"]))
    existing = await _existing_by_phone(session, job.tenant_id, sorted({r.key for r in results if r.key and r.ok}))

    asks_consent = (CONSENT_FIELD in mapping) or options["opt_in_all"]
    consent_blocked = asks_consent and not options["consent_confirmed"]
    counts = {"total": len(results), "valid": 0, "with_warnings": 0, "errors": 0, "duplicate_in_file": 0, "existing_in_db": 0,
              "will_create": 0, "will_update": 0, "will_skip": 0, "will_opt_in": 0, "opted_out_ignored": 0}
    preview: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    for r in results:
        if r.duplicate_in_file_of is not None:
            counts["duplicate_in_file"] += 1
        if not r.ok:
            counts["errors"] += 1
            outcome = "error"
        else:
            counts["valid"] += 1
            if r.warnings:
                counts["with_warnings"] += 1
            match = existing.get(r.key) if r.key else None
            if match is not None:
                counts["existing_in_db"] += 1
                outcome = "update" if options["on_duplicate"] == "update" else "skip"
                counts["will_update" if outcome == "update" else "will_skip"] += 1
            else:
                counts["will_create"] += 1
                outcome = "create"
            if options["consent_confirmed"] and outcome != "skip" and _wants_consent(r.values, options):
                if match is not None and match.opt_out_at is not None:
                    counts["opted_out_ignored"] += 1
                else:
                    counts["will_opt_in"] += 1
        if len(preview) < base.PREVIEW_ROWS:
            preview.append({"row": r.row_number, "outcome": outcome, "values": base._preview_values(r.values),
                            "issues": [i.as_dict() for i in r.issues]})
        for i in r.issues:
            if len(issues) < base.MAX_ISSUES:
                issues.append({"row": r.row_number, **i.as_dict()})

    ignored = sorted(set(table.headers) - set(mapping.values()))
    summary = {"counts": counts, "preview": preview, "issues": issues, "ignored_columns": ignored,
               "required_missing": required_missing, "extracted_by_llm": table.extracted_by_llm,
               "consent": {"asked": asks_consent, "confirmed": options["consent_confirmed"], "blocked": consent_blocked}}
    if consent_blocked:
        # The commit is blocked until the declaration is ticked; surface it like a missing required field.
        summary["required_missing"] = [*required_missing, "إقرار الموافقة على التواصل (PDPL)"]
    job.mapping, job.options, job.summary = mapping, options, summary
    job.total_rows = len(results)
    job.status, job.error_message = "validated", None
    await session.flush()
    return summary


# ── 3. commit + background run ────────────────────────────────────────────────

async def _apply_batch(session, tenant_id: uuid.UUID, job_id: uuid.UUID, batch: list[RowResult], options: dict[str, Any],
                       filename: str, totals: dict[str, int]) -> dict[str, Any]:
    out: dict[str, Any] = {"created": 0, "updated": 0, "skipped": [], "failed": []}
    ready: list[RowResult] = []
    for r in batch:
        if not r.ok:
            out["failed"].append((r, "; ".join(i.message for i in r.errors)))
        else:
            ready.append(r)
    existing = await _existing_by_phone(session, tenant_id, [r.key for r in ready if r.key])
    consent_ok, source, now = options.get("consent_confirmed", False), f"import:{filename}"[:120], datetime.now(timezone.utc)
    extra_tags, on_dup = options.get("tags") or [], options.get("on_duplicate", "skip")

    to_insert: list[tuple[RowResult, dict[str, Any]]] = []
    for r in ready:
        v = r.values
        match = existing.get(r.key)
        if match is not None:
            if on_dup == "skip":
                out["skipped"].append((r, "موجود مسبقًا (نفس رقم الجوال) — تم التخطي"))
                continue
            for f in ("display_name", "email", "city", "notes"):
                if v.get(f):
                    setattr(match, f, v[f])
            tags = list(dict.fromkeys([*(match.tags or []), *(v.get("tags") or []), *extra_tags]))
            if len(tags) > MAX_TAGS:
                tags = tags[:MAX_TAGS]
            match.tags = tags
            if v.get("is_vip"):
                match.is_vip = True
            if match.phone_e164 is None and r.key:
                match.phone_e164 = r.key
            _apply_consent(match, v, options, consent_ok, source, now, totals)
            out["updated"] += 1
            continue
        row = {"customer_id": uuid.uuid4(), "tenant_id": tenant_id, "unified_phone": r.key, "phone_e164": r.key,
               "display_name": v.get("display_name"), "email": v.get("email"), "city": v.get("city"), "notes": v.get("notes"),
               "tags": list(dict.fromkeys([*(v.get("tags") or []), *extra_tags]))[:MAX_TAGS], "is_vip": bool(v.get("is_vip")),
               "source": "import", "channel_user_ids": {}, "marketing_opt_in": False, "is_processing_restricted": False,
               "vcard_state": "STATE_NEW", "engagement_score": 0, "opt_in_source": None, "opt_in_at": None}
        if consent_ok and _wants_consent(v, options):
            row.update(marketing_opt_in=True, opt_in_source=source, opt_in_at=now)
            totals["opted_in"] += 1
        to_insert.append((r, row))

    if to_insert:
        stmt = pg_insert(Customer).values([row for _, row in to_insert]).on_conflict_do_nothing().returning(Customer.customer_id)
        created = {x[0] for x in await session.execute(stmt)}
        for r, row in to_insert:
            if row["customer_id"] in created:
                out["created"] += 1
            else:
                out["skipped"].append((r, "موجود مسبقًا (أُنشئ في الأثناء أو أُعيدت معالجته) — تم التخطي"))
                if row["marketing_opt_in"]:
                    totals["opted_in"] -= 1
    return out


def _apply_consent(c: Customer, values: dict[str, Any], options: dict[str, Any], consent_ok: bool, source: str,
                   now: datetime, totals: dict[str, int]) -> None:
    flag = values.get(CONSENT_FIELD)
    if flag is False:                                   # an explicit "no" in the file is honoured (stricter privacy)
        if customer_svc.apply_marketing_consent(c, False, source, now=now):
            totals["opted_out"] += 1
    elif consent_ok and _wants_consent(values, options):
        if customer_svc.apply_marketing_consent(c, True, source, now=now):
            totals["opted_in"] += 1


async def run_job(tenant_id: uuid.UUID, import_id: uuid.UUID) -> None:
    cfg = get_settings()
    batch_size = max(1, cfg.import_batch_size)
    try:
        for attempt in range(6):
            async with get_tenant_session(tenant_id) as s:
                if await base._claim(s, import_id):
                    job = await base._load_job(s, import_id, (KIND,))
                    rows_key, mapping, options = job.rows_key, dict(job.mapping or {}), dict(job.options or {})
                    filename, done_rows = job.filename, job.processed_rows
                    totals = {"opted_in": 0, "opted_out": 0, **{k: int(v) for k, v in ((job.summary or {}).get("result") or {}).items()}}
                    break
                state = await s.scalar(select(ImportJob.status).where(ImportJob.import_id == import_id))
            if state not in ("validated", "failed") or attempt == 5:
                return
            await asyncio.sleep(0.5)
        table = base._table_from_json(await base.storage.get(rows_key))
        results = list(validate_table(table, mapping, CUSTOMER_SCHEMA, options.get("defaults")))
        report: list[tuple[RowResult, str, str]] = []

        for start in range(done_rows, len(results), batch_size):
            batch = results[start:start + batch_size]
            async with get_tenant_session(tenant_id) as s:
                job = await base._load_job(s, import_id, (KIND,))
                if job.cancel_requested:
                    job.status, job.finished_at = "cancelled", datetime.now(timezone.utc)
                    break
                out = await _apply_batch(s, tenant_id, import_id, batch, options, filename, totals)
                job.processed_rows = start + len(batch)
                job.created_count += out["created"]
                job.updated_count += out["updated"]
                job.skipped_count += len(out["skipped"])
                job.failed_count += len(out["failed"])
                job.summary = {**(job.summary or {}), "result": dict(totals)}
                job.lease_expires_at = datetime.now(timezone.utc) + timedelta(seconds=base.LEASE_SECONDS)
            report += [(r, "فشل", why) for r, why in out["failed"]] + [(r, "تخطّي", why) for r, why in out["skipped"]]
            await asyncio.sleep(0)

        await _finalize(tenant_id, import_id, table.headers, report, options)
    except Exception as exc:  # noqa: BLE001
        logger.error("customer_import_failed", import_id=str(import_id), tenant_id=str(tenant_id), exc_info=True)
        try:
            async with get_tenant_session(tenant_id) as s:
                await s.execute(update(ImportJob).where(ImportJob.import_id == import_id, ImportJob.status == "running")
                                .values(status="failed", finished_at=datetime.now(timezone.utc),
                                        error_message=f"{type(exc).__name__}: {str(exc)[:300]}"))
        except Exception:  # noqa: BLE001
            logger.error("customer_import_fail_mark_failed", import_id=str(import_id), exc_info=True)


async def _finalize(tenant_id: uuid.UUID, import_id: uuid.UUID, headers: list[str], report, options: dict[str, Any]) -> None:
    errors_key = None
    if report:
        errors_key = base._keys(tenant_id, import_id)["errors"]
        await base.storage.put(errors_key, base.build_error_csv(headers, sorted(report, key=lambda e: e[0].row_number)), "text/csv")
    async with get_tenant_session(tenant_id) as s:
        job = await base._load_job(s, import_id, (KIND,))
        if job.status == "running":
            job.status = "completed"
        job.finished_at = job.finished_at or datetime.now(timezone.utc)
        job.errors_key = errors_key or job.errors_key
        job.lease_expires_at = None
        source_key, rows_key = job.source_key, job.rows_key
        job.source_key = job.rows_key = None
        result = (job.summary or {}).get("result") or {}
        await base.write_audit(
            s, tenant_id=tenant_id, actor_user_id=job.created_by, action="customer_import.completed",
            entity_type="import_job", entity_id=str(import_id),
            details={"created": job.created_count, "updated": job.updated_count, "skipped": job.skipped_count, "failed": job.failed_count,
                     "consent_confirmed": bool(options.get("consent_confirmed")), "opted_in": result.get("opted_in", 0),
                     "opted_out": result.get("opted_out", 0), "filename": job.filename})
    await base.storage.delete(source_key)
    await base.storage.delete(rows_key)
    logger.info("customer_import_finished", import_id=str(import_id), tenant_id=str(tenant_id))


def build_excel_template() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "العملاء"
    ws.sheet_view.rightToLeft = True
    ws.append([f.label_ar for f in CUSTOMER_SCHEMA.fields])
    ws.append([f.example for f in CUSTOMER_SCHEMA.fields])
    fill = PatternFill("solid", fgColor="1A2332")
    for i, f in enumerate(CUSTOMER_SCHEMA.fields, start=1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), fill, Alignment(horizontal="center")
        ws.column_dimensions[c.column_letter].width = max(16, len(f.label_ar) + 6)
    guide = wb.create_sheet("تعليمات")
    guide.sheet_view.rightToLeft = True
    guide.append(["الحقل", "مطلوب؟", "ملاحظات"])
    notes = {
        "phone": "يقبل 0501234567 أو 966501234567 أو +966501234567. هو مفتاح التعرّف على العميل؛ المكرر يُعالج حسب خيار «المكرر».",
        "marketing_opt_in": "اتركه فارغًا إن لم تكن لديك موافقة العميل. لا تُسجَّل الموافقة إلا بعد إقرارك في نافذة الاستيراد.",
        "tags": "افصل بين الوسوم بفاصلة، حتى 20 وسمًا.",
    }
    for f in CUSTOMER_SCHEMA.fields:
        guide.append([f.label_ar, "نعم" if f.required else "لا", notes.get(f.name, "")])
    guide.column_dimensions["A"].width, guide.column_dimensions["C"].width = 28, 90
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
