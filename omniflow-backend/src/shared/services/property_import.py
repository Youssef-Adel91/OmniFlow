"""
Property bulk import: upload -> dry-run validation -> background commit.

Flow (see gateway/routers/property_import.py for the HTTP surface):

  upload    sniff the bytes, parse/extract rows, store the original + normalised rows
            in the private imports bucket, suggest a column mapping
  validate  apply the (user-edited) mapping + options to every row WITHOUT writing:
            counts, a 50-row preview, the first issues, duplicate detection vs the DB
  commit    queue a job; `run_job` processes it in batches (default 500)
  progress  counters on the job row (the UI polls it)

Durability: the job is a DB row with a lease. `run_job` claims it atomically
(`queued` or an expired `running` lease), renews the lease every batch, and each
batch commits its rows together with `processed_rows`, so a crash resumes at the
next unprocessed batch without duplicating rows. Rows without a REGA number get
the deterministic placeholder DEV-IMP-<job>-<row>, so re-processing a row is
idempotent (ON CONFLICT DO NOTHING). `resume_stale_jobs()` re-adopts abandoned
jobs when the gateway starts.

Qdrant: each batch's created/updated listings are embedded and upserted in one
call AFTER the batch commit. An outage only leaves those listings un-indexed
(`indexed_count` shows it); it never fails the import.
"""
from __future__ import annotations

import asyncio
import csv
import io
import json
import re
import uuid
from datetime import datetime, timedelta, timezone
from typing import Any

import structlog
from sqlalchemy import String, cast, func, select, text, update
from sqlalchemy.dialects.postgresql import insert as pg_insert

from src.shared.core.config import get_settings
from src.shared.core.exceptions import NotFoundError
from src.shared.db.models import ImportJob, ImportMappingTemplate, PropertyListing
from src.shared.db.session import get_system_session, get_tenant_session
from src.shared.services import vector_sync
from src.shared.services.bulk_import import llm_extract
from src.shared.services.bulk_import.files import (
    LLM_KINDS, ImportFileError, Table, detect_kind, read_tabular,
)
from src.shared.services.bulk_import.mapping import suggest_mapping, validate_mapping
from src.shared.services.bulk_import.property_schema import PROPERTY_SCHEMA
from src.shared.services.bulk_import.validate import RowResult, validate_table

logger = structlog.get_logger(__name__)

KIND = "properties"
LEASE_SECONDS = 120
ON_DUPLICATE = ("skip", "update", "create_new")
DEFAULT_KEYS = ("property_type", "status", "city", "district")
TERMINAL = ("completed", "failed", "cancelled")
ACTIVE = ("queued", "running")
PREVIEW_ROWS = 50
MAX_ISSUES = 100
_RUNNING: set[uuid.UUID] = set()      # job ids this process is working on (avoid double tasks)


# ══════════════════════════════════════════════════════════════════════════════
# Object storage (private bucket); replaceable in tests
# ══════════════════════════════════════════════════════════════════════════════

class ImportStorage:
    def __init__(self) -> None:
        self._bucket_ready = False

    @property
    def bucket(self) -> str:
        return get_settings().s3_imports_bucket

    async def _ensure_bucket(self) -> None:
        if self._bucket_ready:
            return
        from botocore.exceptions import ClientError

        from src.shared.storage.s3 import s3_mgr

        async with s3_mgr.get_client() as client:
            try:
                await client.head_bucket(Bucket=self.bucket)
            except ClientError as exc:
                code = str(exc.response.get("Error", {}).get("Code", ""))
                if code not in ("404", "NoSuchBucket", "NotFound"):
                    raise
                cfg = get_settings()
                kwargs: dict[str, Any] = {"Bucket": self.bucket}
                if not (cfg.minio_endpoint or cfg.s3_endpoint_url) and cfg.aws_region != "us-east-1":
                    kwargs["CreateBucketConfiguration"] = {"LocationConstraint": cfg.aws_region}
                await client.create_bucket(**kwargs)     # private by default (no anonymous policy)
        self._bucket_ready = True

    async def put(self, key: str, data: bytes, content_type: str) -> None:
        from src.shared.storage.s3 import s3_mgr

        await self._ensure_bucket()
        await s3_mgr.upload_file(self.bucket, key, data, content_type)

    async def get(self, key: str) -> bytes:
        from src.shared.storage.s3 import s3_mgr

        return await s3_mgr.download_file(self.bucket, key)

    async def delete(self, key: str | None) -> None:
        if not key:
            return
        from src.shared.storage.s3 import s3_mgr

        try:
            await s3_mgr.delete_file(self.bucket, key)
        except Exception:  # noqa: BLE001 - cleanup is best-effort
            logger.warning("import_storage_delete_failed", key=key, exc_info=True)


storage: Any = ImportStorage()


def _keys(tenant_id: uuid.UUID, import_id: uuid.UUID) -> dict[str, str]:
    base = f"imports/{tenant_id}/{import_id}"
    return {"source": f"{base}/source", "rows": f"{base}/rows.json", "errors": f"{base}/errors.csv"}


# ══════════════════════════════════════════════════════════════════════════════
# Helpers
# ══════════════════════════════════════════════════════════════════════════════

def _safe_filename(name: str) -> str:
    name = (name or "upload").replace("\\", "/").rsplit("/", 1)[-1]
    name = re.sub(r"[\x00-\x1f\x7f]", "", name).strip() or "upload"
    return name[:255]


def csv_safe(value: Any) -> str:
    """Neutralise spreadsheet formula injection (=, +, -, @, tab, CR) in a CSV cell."""
    s = "" if value is None else str(value)
    return "'" + s if s[:1] in ("=", "+", "-", "@", "\t", "\r") else s


def _table_to_json(table: Table) -> bytes:
    return json.dumps(
        {"headers": table.headers, "rows": table.rows, "extracted_by_llm": table.extracted_by_llm, "notes": table.notes},
        ensure_ascii=False,
    ).encode("utf-8")


def _table_from_json(data: bytes) -> Table:
    doc = json.loads(data)
    return Table(doc["headers"], doc["rows"], doc.get("extracted_by_llm", False), doc.get("notes", []))


async def _load_job(session, import_id: uuid.UUID) -> ImportJob:
    job = await session.scalar(select(ImportJob).where(ImportJob.import_id == import_id, ImportJob.kind == KIND))
    if job is None:
        raise NotFoundError("عملية الاستيراد غير موجودة")
    return job


def serialize_job(job: ImportJob) -> dict[str, Any]:
    total = job.total_rows or 0
    pct = 100 if job.status == "completed" else (round(job.processed_rows / total * 100, 1) if total else 0)
    return {
        "import_id": str(job.import_id), "kind": job.kind, "status": job.status, "filename": job.filename,
        "file_kind": job.file_kind, "total_rows": total, "processed_rows": job.processed_rows, "percent": pct,
        "created": job.created_count, "updated": job.updated_count, "skipped": job.skipped_count,
        "failed": job.failed_count, "indexed": job.indexed_count,
        "has_error_report": bool(job.errors_key), "error_message": job.error_message,
        "cancel_requested": job.cancel_requested, "options": job.options, "mapping": job.mapping,
        "summary": job.summary,
        "created_at": job.created_at.isoformat() if job.created_at else None,
        "started_at": job.started_at.isoformat() if job.started_at else None,
        "finished_at": job.finished_at.isoformat() if job.finished_at else None,
    }


# ══════════════════════════════════════════════════════════════════════════════
# 1. Upload
# ══════════════════════════════════════════════════════════════════════════════

async def create_upload(session, *, tenant_id: uuid.UUID, user_id: uuid.UUID | None, filename: str, data: bytes) -> dict[str, Any]:
    cfg = get_settings()
    max_bytes = cfg.import_max_file_mb * 1024 * 1024
    if len(data) > max_bytes:
        raise ImportFileError("file_too_large", f"حجم الملف يتجاوز الحد المسموح ({cfg.import_max_file_mb} ميجابايت).")
    filename = _safe_filename(filename)
    kind = detect_kind(data, filename)

    if kind in LLM_KINDS:
        table = await llm_extract.extract_table(kind, data, PROPERTY_SCHEMA)
    else:
        table = read_tabular(data, kind, cfg.import_max_rows)

    suggestions = suggest_mapping(table.headers, PROPERTY_SCHEMA)
    unmapped_required = [s.field for s in suggestions if not s.header and PROPERTY_SCHEMA.by_name()[s.field].required]
    if unmapped_required and not table.extracted_by_llm:
        taken = {s.header for s in suggestions if s.header}
        refined = await llm_extract.refine_mapping(table.headers, table.rows, unmapped_required, taken, PROPERTY_SCHEMA)
        for s in suggestions:
            if s.field in refined:
                s.header, s.confidence, s.source = refined[s.field], 0.6, "llm"

    import_id = uuid.uuid4()
    keys = _keys(tenant_id, import_id)
    await storage.put(keys["source"], data, "application/octet-stream")
    await storage.put(keys["rows"], _table_to_json(table), "application/json")

    job = ImportJob(
        import_id=import_id, tenant_id=tenant_id, kind=KIND, status="uploaded", filename=filename, file_kind=kind,
        file_size=len(data), source_key=keys["source"], rows_key=keys["rows"], columns=table.headers,
        total_rows=len(table.rows), created_by=user_id,
    )
    session.add(job)
    await session.flush()
    logger.info("import_uploaded", import_id=str(import_id), tenant_id=str(tenant_id), kind=kind, rows=len(table.rows))

    return {
        "import_id": str(import_id), "file_kind": kind, "filename": filename, "total_rows": len(table.rows),
        "detected_columns": table.headers, "suggested_mapping": [s.as_dict() for s in suggestions],
        "fields": [
            {"name": f.name, "label": f.label_ar, "required": f.required, "example": f.example}
            for f in PROPERTY_SCHEMA.fields
        ],
        "sample_rows": [dict(zip(table.headers, r)) for r in table.rows[:5]],
        "extracted_by_llm": table.extracted_by_llm, "notes": table.notes,
        "limits": {"max_file_mb": cfg.import_max_file_mb, "max_rows": cfg.import_max_rows},
    }


# ══════════════════════════════════════════════════════════════════════════════
# 2. Validate (dry run)
# ══════════════════════════════════════════════════════════════════════════════

def normalize_options(options: dict[str, Any] | None) -> dict[str, Any]:
    options = options or {}
    on_dup = options.get("on_duplicate", "skip")
    if on_dup not in ON_DUPLICATE:
        raise ValueError(f"on_duplicate يجب أن يكون أحد: {', '.join(ON_DUPLICATE)}")
    raw_defaults = {k: v for k, v in (options.get("defaults") or {}).items() if v not in (None, "")}
    unknown = set(raw_defaults) - set(DEFAULT_KEYS)
    if unknown:
        raise ValueError(f"قيم افتراضية غير مدعومة: {', '.join(sorted(unknown))}")
    specs = PROPERTY_SCHEMA.by_name()
    from src.shared.services.bulk_import.text import ParseError

    defaults: dict[str, Any] = {}
    for key, value in raw_defaults.items():
        try:
            parsed = specs[key].parse(str(value)) if specs[key].parse else value
        except ParseError as exc:
            raise ValueError(f"قيمة افتراضية غير صالحة — {exc}") from exc
        defaults[key] = str(value) if key in ("city", "district") else parsed
    return {"on_duplicate": on_dup, "defaults": defaults}


async def _existing_regas(session, tenant_id: uuid.UUID, regas: list[str]) -> dict[str, uuid.UUID]:
    found: dict[str, uuid.UUID] = {}
    for i in range(0, len(regas), 1000):
        chunk = regas[i:i + 1000]
        rows = await session.execute(
            select(PropertyListing.rega_ad_number, PropertyListing.listing_id).where(
                PropertyListing.tenant_id == tenant_id, PropertyListing.rega_ad_number.in_(chunk)))
        found.update({r[0]: r[1] for r in rows})
    return found


def _preview_values(values: dict[str, Any]) -> dict[str, Any]:
    return {k: (v if not isinstance(v, str) or len(v) <= 80 else v[:77] + "…") for k, v in values.items()}


async def validate_job(session, job: ImportJob, mapping_in: dict[str, str | None], options_in: dict[str, Any] | None) -> dict[str, Any]:
    if job.status in ACTIVE or job.status in TERMINAL:
        raise ValueError("لا يمكن إعادة التحقق من عملية استيراد بدأت أو انتهت.")
    table = _table_from_json(await storage.get(job.rows_key))
    mapping = validate_mapping(mapping_in, table.headers, PROPERTY_SCHEMA)
    options = normalize_options(options_in)

    required_missing = [
        f.label_ar for f in PROPERTY_SCHEMA.fields
        if f.required and f.name not in mapping and f.name not in options["defaults"]
    ]
    results = list(validate_table(table, mapping, PROPERTY_SCHEMA, options["defaults"]))
    existing = await _existing_regas(session, job.tenant_id, sorted({r.key for r in results if r.key and r.ok}))

    counts = {"total": len(results), "valid": 0, "with_warnings": 0, "errors": 0, "duplicate_in_file": 0,
              "existing_in_db": 0, "will_create": 0, "will_update": 0, "will_skip": 0}
    preview: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    for r in results:
        if r.duplicate_in_file_of is not None:
            counts["duplicate_in_file"] += 1
        if not r.ok:
            counts["errors"] += 1
            outcome = "error"
        else:
            counts["valid"] += 1
            if r.warnings:
                counts["with_warnings"] += 1
            if r.key and r.key in existing:
                counts["existing_in_db"] += 1
                counts[{"skip": "will_skip", "update": "will_update", "create_new": "will_create"}[options["on_duplicate"]]] += 1
                outcome = {"skip": "skip", "update": "update", "create_new": "create"}[options["on_duplicate"]]
            else:
                counts["will_create"] += 1
                outcome = "create"
        if len(preview) < PREVIEW_ROWS:
            preview.append({"row": r.row_number, "outcome": outcome, "values": _preview_values(r.values),
                            "issues": [i.as_dict() for i in r.issues]})
        for i in r.issues:
            if len(issues) < MAX_ISSUES:
                issues.append({"row": r.row_number, **i.as_dict()})

    ignored = sorted(set(table.headers) - set(mapping.values()))
    summary = {"counts": counts, "preview": preview, "issues": issues, "ignored_columns": ignored,
               "required_missing": required_missing, "extracted_by_llm": table.extracted_by_llm}

    job.mapping, job.options, job.summary = mapping, options, summary
    job.total_rows = len(results)
    job.status = "validated"
    job.error_message = None
    await session.flush()
    return summary


# ══════════════════════════════════════════════════════════════════════════════
# 3. Commit + background run
# ══════════════════════════════════════════════════════════════════════════════

class ConflictError(RuntimeError):
    """Another import is already running for this tenant."""


async def queue_job(session, job: ImportJob) -> None:
    if job.status not in ("validated", "failed"):
        raise ValueError("يجب التحقق من الملف (معاينة) قبل الاستيراد.")
    summary = job.summary or {}
    if summary.get("required_missing"):
        raise ValueError("حقول مطلوبة غير مربوطة: " + "، ".join(summary["required_missing"]))
    busy = await session.scalar(
        select(ImportJob.import_id).where(ImportJob.tenant_id == job.tenant_id, ImportJob.status.in_(ACTIVE),
                                          ImportJob.import_id != job.import_id).limit(1))
    if busy is not None:
        raise ConflictError("هناك عملية استيراد أخرى قيد التنفيذ. انتظر حتى تنتهي.")
    job.status, job.cancel_requested, job.error_message = "queued", False, None
    job.lease_expires_at = None
    await session.flush()


def spawn(tenant_id: uuid.UUID, import_id: uuid.UUID) -> None:
    """Start the job on this process's event loop (fire-and-forget; the DB row is the source of truth)."""
    if import_id in _RUNNING:
        return
    _RUNNING.add(import_id)

    async def _go() -> None:
        try:
            await run_job(tenant_id, import_id)
        finally:
            _RUNNING.discard(import_id)

    task = asyncio.create_task(_go(), name=f"import-{import_id}")
    task.add_done_callback(lambda t: t.cancelled() or t.exception())   # retrieve, don't crash the loop


async def start_job(tenant_id: uuid.UUID, import_id: uuid.UUID) -> None:
    """Async entry for Starlette BackgroundTasks (a sync callable would run in a thread, with no event loop)."""
    spawn(tenant_id, import_id)


async def resume_stale_jobs() -> int:
    """Adopt jobs whose worker died (queued, or running with an expired lease). Called at gateway startup."""
    async with get_system_session() as s:
        rows = (await s.execute(text(
            "SELECT import_id, tenant_id FROM import_jobs WHERE status IN ('queued','running') "
            "AND (lease_expires_at IS NULL OR lease_expires_at < now())"))).all()
    for import_id, tenant_id in rows:
        spawn(tenant_id, import_id)
    if rows:
        logger.info("import_jobs_resumed", count=len(rows))
    return len(rows)


async def _claim(session, import_id: uuid.UUID) -> bool:
    now = datetime.now(timezone.utc)
    res = await session.execute(
        update(ImportJob)
        .where(ImportJob.import_id == import_id,
               (ImportJob.status == "queued") | ((ImportJob.status == "running") & (ImportJob.lease_expires_at < now)))
        .values(status="running", lease_expires_at=now + timedelta(seconds=LEASE_SECONDS),
                started_at=func.coalesce(ImportJob.started_at, now)))
    return res.rowcount == 1


def _listing_values(tenant_id: uuid.UUID, values: dict[str, Any], rega: str) -> dict[str, Any]:
    return {
        "listing_id": uuid.uuid4(), "tenant_id": tenant_id, "rega_ad_number": rega,
        "property_type": values["property_type"], "status": values.get("status", "PENDING_VERIFICATION"),
        "is_verified": False,
        **{k: values.get(k) for k in ("city", "district", "latitude", "longitude", "price", "area_sqm",
                                      "bedrooms", "bathrooms", "description_ar", "description_en")},
    }


_UPDATABLE = ("property_type", "status", "city", "district", "latitude", "longitude", "price", "area_sqm",
              "bedrooms", "bathrooms", "description_ar", "description_en")


async def _apply_batch(session, tenant_id: uuid.UUID, import_id: uuid.UUID, batch: list[RowResult], on_dup: str) -> dict[str, Any]:
    """Write one batch. Returns ids to index and per-row skips/failures. Runs inside the caller's transaction."""
    out: dict[str, Any] = {"created": [], "updated": [], "skipped": [], "failed": []}
    ready: list[tuple[RowResult, str]] = []
    for r in batch:
        if not r.ok:
            out["failed"].append((r, "; ".join(i.message for i in r.errors)))
            continue
        ready.append((r, r.key or f"DEV-IMP-{import_id.hex[:8]}-{r.row_number}"))

    existing = await _existing_regas(session, tenant_id, [rega for r, rega in ready if r.key])
    to_insert: list[tuple[RowResult, dict[str, Any]]] = []
    for r, rega in ready:
        if r.key and rega in existing:
            if on_dup == "skip":
                out["skipped"].append((r, "موجود مسبقًا (نفس رقم الإعلان) — تم التخطي"))
                continue
            if on_dup == "update":
                changes = {k: r.values[k] for k in _UPDATABLE if r.values.get(k) is not None}
                if changes:
                    await session.execute(update(PropertyListing).where(
                        PropertyListing.tenant_id == tenant_id, PropertyListing.listing_id == existing[rega]).values(**changes))
                out["updated"].append(existing[rega])
                continue
            rega = f"DEV-IMP-{import_id.hex[:8]}-{r.row_number}"       # create_new: keep both listings
        to_insert.append((r, _listing_values(tenant_id, r.values, rega)))

    if to_insert:
        stmt = (pg_insert(PropertyListing).values([v for _, v in to_insert])
                .on_conflict_do_nothing(constraint="uq_listings_tenant_rega")
                .returning(PropertyListing.listing_id))
        created = {row[0] for row in await session.execute(stmt)}
        for r, v in to_insert:
            if v["listing_id"] in created:
                out["created"].append(v["listing_id"])
            else:
                out["skipped"].append((r, "موجود مسبقًا (أُنشئ في الأثناء أو أُعيدت معالجته) — تم التخطي"))
    return out


async def _index(tenant_id: uuid.UUID, ids: list[uuid.UUID]) -> int:
    """Embed + upsert listings, then record qdrant_point_id. Never raises."""
    if not ids:
        return 0
    try:
        async with get_tenant_session(tenant_id) as s:
            cols = PropertyListing.__table__.c
            rows = (await s.execute(select(*cols).where(PropertyListing.tenant_id == tenant_id,
                                                         PropertyListing.listing_id.in_(ids)))).mappings().all()
        listings = [PropertyListing(**dict(r)) for r in rows]
        done = await vector_sync.sync_listings_batch(listings, tenant_id)
        if done:
            async with get_tenant_session(tenant_id) as s:
                await s.execute(update(PropertyListing).where(
                    PropertyListing.tenant_id == tenant_id, PropertyListing.listing_id.in_([uuid.UUID(i) for i in done]))
                    .values(qdrant_point_id=cast(PropertyListing.listing_id, String)))
        return len(done)
    except Exception:  # noqa: BLE001
        logger.error("import_index_failed", tenant_id=str(tenant_id), exc_info=True)
        return 0


async def reindex(tenant_id: uuid.UUID, ids: list[uuid.UUID]) -> int:
    """Re-embed listings after a bulk edit (background). Public alias of the importer's indexer."""
    return await _index(tenant_id, ids)


def build_error_csv(headers: list[str], entries: list[tuple[RowResult, str, str]]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["رقم الصف", "الحالة", "السبب", *headers])
    for r, status, reason in entries:
        w.writerow([r.row_number, status, csv_safe(reason), *(csv_safe(r.raw.get(h, "")) for h in headers)])
    return b"\xef\xbb\xbf" + buf.getvalue().encode("utf-8")


async def run_job(tenant_id: uuid.UUID, import_id: uuid.UUID) -> None:
    cfg = get_settings()
    batch_size = max(1, cfg.import_batch_size)
    try:
        for attempt in range(6):
            async with get_tenant_session(tenant_id) as s:
                if await _claim(s, import_id):
                    job = await _load_job(s, import_id)
                    rows_key, mapping, options = job.rows_key, dict(job.mapping or {}), dict(job.options or {})
                    done_rows = job.processed_rows
                    break
                state = await s.scalar(select(ImportJob.status).where(ImportJob.import_id == import_id))
            if state not in ("validated", "failed") or attempt == 5:
                return                      # owned by another worker, finished, or not runnable
            await asyncio.sleep(0.5)        # the commit that queued it may not be visible yet
        table = _table_from_json(await storage.get(rows_key))
        results = list(validate_table(table, mapping, PROPERTY_SCHEMA, options.get("defaults")))
        on_dup = options.get("on_duplicate", "skip")
        report: list[tuple[RowResult, str, str]] = []

        for start in range(done_rows, len(results), batch_size):
            batch = results[start:start + batch_size]
            async with get_tenant_session(tenant_id) as s:
                job = await _load_job(s, import_id)
                if job.cancel_requested:
                    job.status, job.finished_at = "cancelled", datetime.now(timezone.utc)
                    break
                out = await _apply_batch(s, tenant_id, import_id, batch, on_dup)
                job.processed_rows = start + len(batch)
                job.created_count += len(out["created"])
                job.updated_count += len(out["updated"])
                job.skipped_count += len(out["skipped"])
                job.failed_count += len(out["failed"])
                job.lease_expires_at = datetime.now(timezone.utc) + timedelta(seconds=LEASE_SECONDS)
            report += [(r, "فشل", why) for r, why in out["failed"]] + [(r, "تخطّي", why) for r, why in out["skipped"]]
            indexed = await _index(tenant_id, out["created"] + out["updated"])
            if indexed:
                async with get_tenant_session(tenant_id) as s:
                    await s.execute(update(ImportJob).where(ImportJob.import_id == import_id)
                                    .values(indexed_count=ImportJob.indexed_count + indexed))
            await asyncio.sleep(0)                      # let the API breathe between batches

        await _finalize(tenant_id, import_id, table.headers, report)
    except Exception as exc:  # noqa: BLE001
        logger.error("import_job_failed", import_id=str(import_id), tenant_id=str(tenant_id), exc_info=True)
        try:
            async with get_tenant_session(tenant_id) as s:
                await s.execute(update(ImportJob).where(ImportJob.import_id == import_id, ImportJob.status == "running")
                                .values(status="failed", finished_at=datetime.now(timezone.utc),
                                        error_message=f"{type(exc).__name__}: {str(exc)[:300]}"))
        except Exception:  # noqa: BLE001
            logger.error("import_job_fail_mark_failed", import_id=str(import_id), exc_info=True)


async def _finalize(tenant_id: uuid.UUID, import_id: uuid.UUID, headers: list[str], report: list[tuple[RowResult, str, str]]) -> None:
    errors_key = None
    if report:
        errors_key = _keys(tenant_id, import_id)["errors"]
        await storage.put(errors_key, build_error_csv(headers, sorted(report, key=lambda e: e[0].row_number)), "text/csv")
    async with get_tenant_session(tenant_id) as s:
        job = await _load_job(s, import_id)
        if job.status == "running":
            job.status = "completed"
        job.finished_at = job.finished_at or datetime.now(timezone.utc)
        job.errors_key = errors_key or job.errors_key
        job.lease_expires_at = None
        source_key, rows_key = job.source_key, job.rows_key
        job.source_key = job.rows_key = None
    # The source sheet holds customer/listing data: don't keep it once the job is over.
    await storage.delete(source_key)
    await storage.delete(rows_key)
    logger.info("import_job_finished", import_id=str(import_id), tenant_id=str(tenant_id))


async def cancel_job(session, job: ImportJob) -> None:
    if job.status in TERMINAL:
        raise ValueError("العملية انتهت بالفعل.")
    if job.status in ("uploaded", "validated"):
        job.status, job.finished_at = "cancelled", datetime.now(timezone.utc)
        await storage.delete(job.source_key)
        await storage.delete(job.rows_key)
        job.source_key = job.rows_key = None
    else:
        job.cancel_requested = True            # the running loop stops before its next batch


async def error_report(job: ImportJob) -> bytes:
    if not job.errors_key:
        raise NotFoundError("لا يوجد تقرير أخطاء لهذه العملية")
    return await storage.get(job.errors_key)


# ══════════════════════════════════════════════════════════════════════════════
# Mapping templates + the downloadable Excel template
# ══════════════════════════════════════════════════════════════════════════════

async def list_templates(session) -> list[dict[str, Any]]:
    rows = (await session.execute(select(ImportMappingTemplate).where(ImportMappingTemplate.kind == KIND)
                                  .order_by(ImportMappingTemplate.name))).scalars().all()
    return [{"template_id": str(t.template_id), "name": t.name, "mapping": t.mapping} for t in rows]


async def save_template(session, tenant_id: uuid.UUID, name: str, mapping: dict[str, str]) -> dict[str, Any]:
    name = name.strip()[:100]
    if not name:
        raise ValueError("اسم القالب مطلوب")
    clean = {k: v for k, v in mapping.items() if v and k in PROPERTY_SCHEMA.by_name()}
    existing = await session.scalar(select(ImportMappingTemplate).where(
        ImportMappingTemplate.kind == KIND, ImportMappingTemplate.name == name))
    if existing:
        existing.mapping = clean
        tpl = existing
    else:
        tpl = ImportMappingTemplate(tenant_id=tenant_id, kind=KIND, name=name, mapping=clean)
        session.add(tpl)
    await session.flush()
    return {"template_id": str(tpl.template_id), "name": tpl.name, "mapping": tpl.mapping}


async def delete_template(session, template_id: uuid.UUID) -> None:
    tpl = await session.scalar(select(ImportMappingTemplate).where(
        ImportMappingTemplate.template_id == template_id, ImportMappingTemplate.kind == KIND))
    if tpl is None:
        raise NotFoundError("القالب غير موجود")
    await session.delete(tpl)


def build_excel_template() -> bytes:
    """Arabic-header workbook with one example row and an instructions sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    plain = [f for f in PROPERTY_SCHEMA.fields if f.name != "coordinates"]
    wb = Workbook()
    ws = wb.active
    ws.title = "العقارات"
    ws.sheet_view.rightToLeft = True
    ws.append([f.label_ar for f in plain])
    ws.append([f.example for f in plain])
    head_fill = PatternFill("solid", fgColor="1A2332")
    for i, f in enumerate(plain, start=1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), head_fill, Alignment(horizontal="center")
        ws.column_dimensions[c.column_letter].width = max(16, len(f.label_ar) + 6)
    guide = wb.create_sheet("تعليمات")
    guide.sheet_view.rightToLeft = True
    guide.append(["الحقل", "مطلوب؟", "ملاحظات"])
    notes = {
        "property_type": "القيم: " + "، ".join(("شقة", "فيلا", "أرض", "تجاري", "إيجار يومي", "مكتب", "مستودع")),
        "status": "اتركه فارغًا ليصبح «قيد التحقق». القيم: " + "، ".join(("قيد التحقق", "نشط ومعتمد", "موقوف", "مباع", "مؤجر", "مسحوب")),
        "price": "بالريال. يقبل «850000» أو «1.2 مليون» أو «٨٥٠ ألف».",
        "area_sqm": "بالمتر المربع. يقبل «180» أو «180 م²».",
        "rega_ad_number": "اتركه فارغًا لإنشاء رقم مؤقت (DEV-IMP). الأرقام المكررة تُعالج حسب خيار «المكرر».",
    }
    for f in plain:
        guide.append([f.label_ar, "نعم" if f.required else "لا", notes.get(f.name, "")])
    guide.column_dimensions["A"].width, guide.column_dimensions["C"].width = 28, 90
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
